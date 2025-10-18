"""
SlideParser module for extracting content from PowerPoint presentations.
"""

from typing import List, Dict, Any, Optional
from pptx import Presentation
from pptx.util import Inches
import io
from PIL import Image


class SlideContent:
    """Represents the content of a single slide."""

    def __init__(self, slide_number: int):
        self.slide_number = slide_number
        self.title = ""
        self.text_content = []
        self.notes = ""
        self.images_count = 0
        self.has_chart = False
        self.has_table = False

    def to_dict(self) -> Dict[str, Any]:
        """Convert slide content to dictionary format."""
        return {
            "slide_number": self.slide_number,
            "title": self.title,
            "text_content": self.text_content,
            "notes": self.notes,
            "images_count": self.images_count,
            "has_chart": self.has_chart,
            "has_table": self.has_table
        }

    def __str__(self) -> str:
        """String representation of slide content."""
        text = f"Slide {self.slide_number}"
        if self.title:
            text += f": {self.title}"
        text += "\n"

        if self.text_content:
            text += "Content:\n"
            for item in self.text_content:
                text += f"  - {item}\n"

        if self.notes:
            text += f"Notes: {self.notes}\n"

        if self.images_count > 0:
            text += f"Images: {self.images_count}\n"

        if self.has_chart:
            text += "Contains chart\n"

        if self.has_table:
            text += "Contains table\n"

        return text


class SlideParser:
    """Parser for extracting content from PowerPoint presentations."""

    def __init__(self, pptx_path: str):
        """
        Initialize the parser with a PowerPoint file.

        Args:
            pptx_path: Path to the PowerPoint file
        """
        self.pptx_path = pptx_path
        self.presentation = None
        self.slides_content: List[SlideContent] = []

    def load_presentation(self) -> None:
        """Load the PowerPoint presentation."""
        try:
            self.presentation = Presentation(self.pptx_path)
        except Exception as e:
            raise Exception(f"Failed to load presentation: {str(e)}")

    def _extract_text_from_shape(self, shape) -> Optional[str]:
        """Extract text from a shape if it has text."""
        if hasattr(shape, "text") and shape.text.strip():
            return shape.text.strip()
        return None

    def _parse_slide(self, slide, slide_number: int) -> SlideContent:
        """
        Parse a single slide and extract its content.

        Args:
            slide: The slide object from python-pptx
            slide_number: The slide number (1-indexed)

        Returns:
            SlideContent object with extracted information
        """
        content = SlideContent(slide_number)

        # Extract title and content from shapes
        for shape in slide.shapes:
            # Check if shape has text
            text = self._extract_text_from_shape(shape)

            if text:
                # First text frame is typically the title
                if not content.title and hasattr(shape, "placeholder_format"):
                    if shape.placeholder_format.type == 1:  # Title placeholder
                        content.title = text
                        continue

                # If we already have a title, or this isn't a title, add to content
                if text != content.title:
                    content.text_content.append(text)

            # Check for images
            if shape.shape_type == 13:  # Picture type
                content.images_count += 1

            # Check for charts
            if hasattr(shape, "chart"):
                content.has_chart = True

            # Check for tables
            if hasattr(shape, "table"):
                content.has_table = True

        # Extract notes
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            if notes_slide.notes_text_frame:
                content.notes = notes_slide.notes_text_frame.text.strip()

        return content

    def parse_all_slides(self) -> List[SlideContent]:
        """
        Parse all slides in the presentation.

        Returns:
            List of SlideContent objects
        """
        if not self.presentation:
            self.load_presentation()

        self.slides_content = []

        for idx, slide in enumerate(self.presentation.slides, start=1):
            content = self._parse_slide(slide, idx)
            self.slides_content.append(content)

        return self.slides_content

    def get_slide_count(self) -> int:
        """Get the total number of slides in the presentation."""
        if not self.presentation:
            self.load_presentation()
        return len(self.presentation.slides)

    def get_presentation_summary(self) -> str:
        """Get a summary of the presentation structure."""
        if not self.presentation:
            self.load_presentation()

        summary = f"Presentation: {self.pptx_path}\n"
        summary += f"Total slides: {len(self.presentation.slides)}\n"
        summary += f"Slide dimensions: {self.presentation.slide_width} x {self.presentation.slide_height}\n"

        return summary
